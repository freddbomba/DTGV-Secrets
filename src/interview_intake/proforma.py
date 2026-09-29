"""Pro forma interview data: load it and export it to Excel.

The pro forma is a lightweight, spreadsheet-friendly view of interview
metadata, used for tracking and reporting.  It is intentionally pseudonymous:
it carries no real names (the identity mapping stays in the supervisor's
masterfile).

The Excel export uses the optional ``openpyxl`` dependency (extra ``excel``):

    pip install -e '.[excel]'
    pro-forma-to-excel -i tests/data/pro_forma_interviews.json -o pro_forma.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Sequence

# (record key, spreadsheet header)
COLUMNS: tuple[tuple[str, str], ...] = (
    ("interview_id", "Interview ID"),
    ("date", "Date"),
    ("researcher_id", "Researcher"),
    ("mnemonic", "Mnemonic"),
    ("location", "Location"),
    ("language", "Language"),
    ("participants", "Participants"),
    ("duration_seconds", "Duration (s)"),
    ("original_filename", "Source file"),
    ("format", "Format"),
    ("size_bytes", "Size (bytes)"),
    ("transcription_status", "Transcription"),
    ("consent_ref", "Consent ref"),
    ("notes", "Notes"),
)

_REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = _REPO_ROOT / "tests" / "data" / "pro_forma_interviews.json"
DEFAULT_OUTPUT = Path("pro_forma_interviews.xlsx")
DEFAULT_SHEET = "Pro forma"

_NUMBER_COLUMNS = {"participants", "duration_seconds", "size_bytes"}


def load_pro_forma(path: Path | str = DEFAULT_INPUT) -> dict[str, Any]:
    """Load and minimally validate a pro forma JSON document."""
    path = Path(path)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"Pro forma file not found: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"Malformed JSON in {path}: {exc}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("interviews"), list):
        raise ValueError(
            f"{path} must be an object with an 'interviews' list."
        )
    for index, record in enumerate(data["interviews"]):
        if not isinstance(record, Mapping) or "interview_id" not in record:
            raise ValueError(
                f"interviews[{index}] must be an object with an 'interview_id'."
            )
    return data


def rows_for(interviews: Iterable[Mapping[str, Any]]) -> list[list[Any]]:
    """Return a header row followed by one row per interview."""
    rows: list[list[Any]] = [[label for _, label in COLUMNS]]
    for record in interviews:
        rows.append([record.get(key) for key, _ in COLUMNS])
    return rows


def write_excel(
    interviews: Iterable[Mapping[str, Any]],
    output: Path | str = DEFAULT_OUTPUT,
    *,
    sheet: str = DEFAULT_SHEET,
) -> Path:
    """Write the interviews to an ``.xlsx`` workbook and return its path.

    Raises :class:`RuntimeError` if ``openpyxl`` is not installed.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - depends on environment
        raise RuntimeError(
            "Excel export needs openpyxl. Install it with "
            "`pip install -e '.[excel]'` or `pip install openpyxl`."
        ) from exc

    interviews = list(interviews)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet[:31]  # Excel sheet-name limit

    header = [label for _, label in COLUMNS]
    worksheet.append(header)
    for record in interviews:
        worksheet.append([record.get(key) for key, _ in COLUMNS])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for index, (key, label) in enumerate(COLUMNS, start=1):
        widest = len(label)
        for record in interviews:
            widest = max(widest, len(str(record.get(key, ""))))
        worksheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, 60)
        if key in _NUMBER_COLUMNS:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row, column=index).number_format = "#,##0"

    worksheet.freeze_panes = "A2"
    if interviews:
        worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(output)
    return output


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="pro-forma-to-excel",
        description="Convert a pro forma JSON file into an Excel workbook.",
    )
    parser.add_argument(
        "-i", "--input", default=str(DEFAULT_INPUT),
        help=f"Pro forma JSON (default: {DEFAULT_INPUT}).",
    )
    parser.add_argument(
        "-o", "--output", default=str(DEFAULT_OUTPUT),
        help=f"Output .xlsx path (default: {DEFAULT_OUTPUT}).",
    )
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Worksheet name.")
    args = parser.parse_args(argv)

    try:
        data = load_pro_forma(args.input)
        path = write_excel(data["interviews"], args.output, sheet=args.sheet)
    except (ValueError, RuntimeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    print(f"Wrote {path} ({len(data['interviews'])} interviews).")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
